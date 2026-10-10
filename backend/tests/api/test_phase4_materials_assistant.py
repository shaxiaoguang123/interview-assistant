import io,json
from datetime import datetime,timezone
from sqlalchemy import select,func,text
from app.models.material import Project,Material,MaterialVersion,MaterialChunk
from app.models.assistant import AssistantOutput
from app.models.saved_answer import SavedAnswer,SavedAnswerVersion

class FakeProvider:
    provider='fake_test'
    model='fake-local'
    def __init__(self):self.messages=[]
    def complete(self,messages):
        self.messages.append(messages)
        return '先说明设计目标，再解释 LangGraph 持久化和 MCP 通信协议。[S1]'


def upload(client,content='LangGraph 持久化\nMCP 通信协议\nFunction Calling\nRAG 检索重排',filename='resume.md',**metadata):
    response=client.post('/api/v1/materials',data={'file':(io.BytesIO(content.encode()),filename),'metadata':json.dumps(metadata)},content_type='multipart/form-data')
    assert response.status_code==201,response.get_json()
    return response.get_json()


def project(client,**values):
    r=client.post('/api/v1/projects',json={'name':'Agent 项目','summary':'LangGraph 持久化','personal_role':'设计 MCP 通信协议','notes':'PRIVATE MANAGEMENT NOTE',**values})
    assert r.status_code==201,r.get_json();return r.get_json()


def generate(client,fake,qid,operation='reference_answer',**extra):
    payload={'question_id':qid,'operation':operation,'context':{},**extra}
    preview=client.post('/api/v1/assistant/context-preview',json=payload)
    assert preview.status_code==200,preview.get_json()
    result=client.post('/api/v1/assistant/'+operation.replace('_','-'),json={**payload,'context_token':preview.get_json()['context_token']})
    assert result.status_code==200,result.get_json();return result.get_json()


def question(client):
    r=client.post('/api/v1/questions',json={'text':'解释 LangGraph 持久化与 MCP 通信协议'});assert r.status_code==201;return r.get_json()['id']


def test_profile_unique_versioned_facts_and_notes_excluded(client,db_session):
    p=project(client)
    assert p['profile_version_count']==1
    profile=p['profile_material_id']
    v1=client.get(f'/api/v1/materials/{profile}/versions').get_json()[0]
    text1=client.get(f"/api/v1/material-versions/{v1['id']}").get_json()['text']
    assert 'PRIVATE MANAGEMENT NOTE' not in text1
    r=client.patch(f"/api/v1/projects/{p['id']}",json={'summary':'RAG 检索重排与可恢复工具调用'})
    assert r.status_code==200 and r.get_json()['profile_version_count']==2
    assert client.get(f"/api/v1/material-versions/{v1['id']}").get_json()['text']==text1
    assert client.patch(f"/api/v1/projects/{p['id']}",json={'notes':'only private notes','is_active':False}).get_json()['profile_version_count']==2
    assert db_session.scalar(select(func.count()).select_from(Material).where(Material.project_id==p['id'],Material.is_system_managed.is_(True)))==1
    assert client.patch(f'/api/v1/materials/{profile}',json={'title':'manually changed'}).status_code==409
    assert client.post(f'/api/v1/materials/{profile}/versions',data={'file':(io.BytesIO(b'new'),'new.txt')}).status_code==409


def test_upload_versions_safe_paths_and_keyword_search(client,app):
    m=upload(client,filename='../../resume.md',kind='resume')
    v1=m['current_version'];old=client.get(f"/api/v1/material-versions/{v1['id']}").get_json()
    result=client.post(f"/api/v1/materials/{m['id']}/versions",data={'file':(io.BytesIO('新的项目内容'.encode()),'next.txt')})
    assert result.status_code==201 and result.get_json()['version_count']==2
    assert client.get(f"/api/v1/material-versions/{v1['id']}").get_json()['text']==old['text']
    assert client.get(v1['download_url']).data.startswith(b'LangGraph')
    with app.extensions['sqlalchemy_session_factory']() as s:
        rows=list(s.scalars(select(MaterialVersion)))
        assert len({v.path for v in rows})==2 and all('/' not in v.path and '..' not in v.path for v in rows)
    # Search current versions only: old version is retained but never searched by default.
    for query in ('LangGraph 持久化','MCP 通信协议','Function Calling','RAG 检索重排'):
        current=upload(client,kind='resume')
        r=client.post('/api/v1/materials/search',json={'query':query,'context':{'material_ids':[current['id']]}})
        assert r.status_code==200 and r.get_json()['results']
        assert {i['material_version_id'] for i in r.get_json()['results']}=={current['current_version']['id']}
    r=client.post('/api/v1/materials',data={'file':(io.BytesIO(b'invalid pdf'), 'invalid.pdf')})
    assert r.status_code==400 and r.get_json()['error']['code']=='MATERIAL_PARSE_FAILED'
    assert client.post('/api/v1/materials',data={'file':(io.BytesIO(b''),'empty.txt')}).status_code==400


def test_context_disabled_none_and_stale_never_send_private_files(client,app,db_session):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client);p=project(client)
    disabled=upload(client,content='SECRET DISABLED BODY',project_id=p['id'])
    client.patch(f"/api/v1/materials/{disabled['id']}",json={'include_in_context':False})
    generate(client,fake,qid,context={'project_ids':[p['id']]})
    sent=json.dumps(fake.messages[-1],ensure_ascii=False)
    assert 'SECRET DISABLED BODY' not in sent and 'PRIVATE MANAGEMENT NOTE' not in sent
    assert 'LangGraph 持久化' in sent
    unrelated=upload(client,content='UNSELECTED FILE CONTENT')
    generate(client,fake,qid)
    sent=json.loads(fake.messages[-1][1]['content'])
    assert sent['evidence']==[] and 'UNSELECTED FILE CONTENT' not in json.dumps(sent)
    payload={'question_id':qid,'operation':'reference_answer','context':{'material_ids':[unrelated['id']]}}
    preview=client.post('/api/v1/assistant/context-preview',json=payload).get_json()
    client.patch(f"/api/v1/materials/{unrelated['id']}",json={'include_in_context':False})
    count=len(fake.messages)
    assert client.post('/api/v1/assistant/reference-answer',json={**payload,'context_token':preview['context_token']}).status_code==409
    assert len(fake.messages)==count
    assert db_session.scalar(select(func.count()).select_from(AssistantOutput))==0
    assert db_session.scalar(select(func.count()).select_from(SavedAnswer))==0


def test_explicit_polish_save_appends_preserves_mastery_rating_and_old_evidence(client,app,db_session):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client);p=project(client)
    answer=client.post(f'/api/v1/questions/{qid}/saved-answers',json={'content':'我的原始回答','self_rating':5}).get_json()
    session=client.post('/api/v1/practice-sessions',json={'mode':'random','limit':1}).get_json()
    event=client.post(f"/api/v1/session-items/{session['items'][0]['id']}/review",json={'review_rating':'basic'}).get_json()
    state_before=client.get(f'/api/v1/questions/{qid}').get_json()['state']
    generated=generate(client,fake,qid,'polish',source_saved_answer_version_id=answer['current_version']['id'],context={'project_ids':[p['id']]})
    assert client.get(f'/api/v1/questions/{qid}/assistant-outputs').get_json()==[]
    assert len(client.get(f"/api/v1/saved-answers/{answer['id']}/versions").get_json())==1
    old_source=generated['sources'][0]['material_version_id']
    client.patch(f"/api/v1/projects/{p['id']}",json={'summary':'Updated facts after generation'})
    saved=client.post(f"/api/v1/assistant/previews/{generated['preview_id']}/save",json={'save_kind':'answer_version','saved_answer_id':answer['id']})
    assert saved.status_code==201,saved.get_json()
    result=saved.get_json();v=result['version']
    assert v['origin_kind']=='ai_assisted' and v['based_on_version_id']==answer['current_version']['id']
    assert v['assistant_output_id']==result['output']['id'] and v['self_rating'] is None
    all_versions=client.get(f"/api/v1/saved-answers/{answer['id']}/versions").get_json()
    assert len(all_versions)==2 and all_versions[1]['content']=='我的原始回答' and all_versions[1]['self_rating']==5
    assert client.get(f'/api/v1/questions/{qid}').get_json()['state']==state_before
    review=client.get(f'/api/v1/questions/{qid}/practice-reviews').get_json()[0]
    assert review['id']==event['id'] and review['review_rating']=='basic' and review['reviewed_at']==event['reviewed_at']
    assert result['output']['sources'][0]['material_version_id']==old_source
    record=db_session.get(AssistantOutput,result['output']['id'])
    assert '项目' not in json.dumps(record.selected_context_json) and 'text' not in json.dumps(record.selected_context_json)
    assert client.post(f"/api/v1/assistant/previews/{generated['preview_id']}/save",json={'save_kind':'new_answer'}).status_code==409
    assert client.get(f'/api/v1/questions/{qid}/history').get_json()['assistant_outputs'][0]['id']==record.id


def test_reference_origins_analysis_independence_and_archive(client,app):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client)
    exact=generate(client,fake,qid)
    r=client.post(f"/api/v1/assistant/previews/{exact['preview_id']}/save",json={'save_kind':'new_answer'})
    assert r.status_code==201 and r.get_json()['version']['origin_kind']=='ai_generated'
    rewritten=generate(client,fake,qid)
    r=client.post(f"/api/v1/assistant/previews/{rewritten['preview_id']}/save",json={'save_kind':'new_answer','content':'用户修改了参考答案'})
    assert r.get_json()['version']['origin_kind']=='ai_assisted'
    analyzed=generate(client,fake,qid,'analyze',draft='需要分析的回答')
    assert client.post(f"/api/v1/assistant/previews/{analyzed['preview_id']}/save",json={'save_kind':'new_answer'}).status_code==400
    assert client.post(f"/api/v1/assistant/previews/{analyzed['preview_id']}/save",json={'save_kind':'output'}).status_code==201
    assert client.get(f'/api/v1/questions/{qid}/practice-reviews').get_json()==[]
    p=project(client);upload(client,project_id=p['id'])
    client.post(f"/api/v1/projects/{p['id']}/archive")
    assert not client.get('/api/v1/projects').get_json()
    assert client.get('/api/v1/projects?include_archived=1').get_json()
    assert client.post('/api/v1/assistant/context-preview',json={'question_id':qid,'operation':'reference_answer','context':{'project_ids':[p['id']]}}).status_code==409


def test_pdf_docx_parsing_and_scanned_pdf_not_success(client):
    from pypdf import PdfWriter
    from docx import Document
    pdf=PdfWriter();pdf.add_blank_page(width=600,height=800);binary=io.BytesIO();pdf.write(binary)
    r=client.post('/api/v1/materials',data={'file':(io.BytesIO(binary.getvalue()),'scan.pdf')})
    assert r.status_code==400 and r.get_json()['error']['code']=='MATERIAL_NO_TEXT'
    # Minimal valid text PDF, no external renderer or real personal file.
    from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
    page=pdf.pages[0];font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):pdf._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 14 Tf 50 700 Td (MCP Function Calling Evidence) Tj ET')
    page[NameObject('/Contents')]=pdf._add_object(stream);binary=io.BytesIO();pdf.write(binary)
    r=client.post('/api/v1/materials',data={'file':(io.BytesIO(binary.getvalue()),'text.pdf')})
    assert r.status_code==201,r.get_json()
    assert 'MCP' in client.get(f"/api/v1/material-versions/{r.get_json()['current_version']['id']}").get_json()['text']
    document=Document();document.add_paragraph('LangGraph 持久化与 checkpoint');table=document.add_table(rows=1,cols=2);table.cell(0,0).text='Role';table.cell(0,1).text='Developer'
    binary=io.BytesIO();document.save(binary)
    r=client.post('/api/v1/materials',data={'file':(io.BytesIO(binary.getvalue()),'resume.docx')})
    assert r.status_code==201,r.get_json()
    body=client.get(f"/api/v1/material-versions/{r.get_json()['current_version']['id']}").get_json()['text']
    assert 'LangGraph' in body and 'Developer' in body


def test_project_transaction_failure_leaves_no_files_or_profile(client,app,monkeypatch):
    from pathlib import Path
    def fail(*args,**kwargs):raise RuntimeError('injected chunk failure')
    monkeypatch.setattr('app.services.materials.append_material_version',fail)
    r=client.post('/api/v1/projects',json={'name':'Fail project'})
    assert r.status_code==500
    with app.extensions['sqlalchemy_session_factory']() as s:
        assert s.scalar(select(func.count()).select_from(Project))==0
        assert s.scalar(select(func.count()).select_from(Material))==0
    assert not list((Path(app.config['APP_DATA_DIR'])/'materials').glob('*'))


def test_file_write_then_chunk_failure_rolls_back_and_cleans(client,app,monkeypatch):
    from pathlib import Path
    from sqlalchemy import event
    engine=app.extensions['sqlalchemy_engine']
    def fail(conn,cursor,statement,params,context,many):
        if statement.startswith('INSERT INTO material_chunk '):raise RuntimeError('injected chunk persistence failure')
    event.listen(engine,'before_cursor_execute',fail)
    try:
        r=client.post('/api/v1/materials',data={'file':(io.BytesIO(b'MCP evidence'),'file.txt')})
        assert r.status_code==500
    finally:event.remove(engine,'before_cursor_execute',fail)
    with app.extensions['sqlalchemy_session_factory']() as s:
        assert s.scalar(select(func.count()).select_from(Material))==0
        assert s.scalar(select(func.count()).select_from(MaterialVersion))==0
    assert not list((Path(app.config['APP_DATA_DIR'])/'materials').glob('*'))


def test_explicit_saved_reference_can_later_become_answer_and_keeps_provenance(client,app):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client)
    result=generate(client,fake,qid)
    saved=client.post(f"/api/v1/assistant/previews/{result['preview_id']}/save",json={'save_kind':'output'}).get_json()
    assert client.get(f'/api/v1/questions/{qid}/saved-answers').get_json()==[]
    response=client.post(f"/api/v1/assistant-outputs/{saved['output']['id']}/saved-answer",json={'save_kind':'new_answer','content':'我根据参考答案整理的表达'})
    assert response.status_code==201,response.get_json()
    answer=response.get_json()['answer'];current=answer['current_version']
    assert current['origin_kind']=='ai_assisted'
    appended=client.post(f"/api/v1/saved-answers/{answer['id']}/versions",json={'content':'进一步完善表达'})
    assert appended.status_code==201 and appended.get_json()['origin_kind']=='ai_assisted'
    assert appended.get_json()['based_on_version_id']==current['id']
    assert appended.get_json()['assistant_output_id']==saved['output']['id']


def test_archive_keeps_sources_and_dependency_delete_guard(client,app):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client);m=upload(client)
    preview=generate(client,fake,qid,context={'material_ids':[m['id']]})
    client.post(f"/api/v1/assistant/previews/{preview['preview_id']}/save",json={'save_kind':'output'})
    assert client.get(f"/api/v1/materials/{m['id']}/deletion-impact").get_json()['can_delete'] is False
    assert client.delete(f"/api/v1/materials/{m['id']}",json={'confirm':True}).status_code==409
    assert client.post(f"/api/v1/materials/{m['id']}/archive").status_code==200
    assert client.get('/api/v1/materials').get_json()==[]
    assert client.get(f"/api/v1/material-versions/{m['current_version']['id']}").status_code==200
    clean=upload(client,content='No dependencies')
    assert client.delete(f"/api/v1/materials/{clean['id']}",json={'confirm':True}).status_code==200
    assert client.get(f"/api/v1/material-versions/{clean['current_version']['id']}").status_code==404


def test_config_and_provider_failures_never_return_secrets(client,app,monkeypatch):
    import urllib.error
    from app.services.llm_provider import CompatibleProvider
    r=client.patch('/api/v1/llm/config',json={'base_url':'http://127.0.0.1:12345/v1','model':'test-model','api_key':'test-secret'})
    assert r.status_code==200 and 'test-secret' not in r.get_data(as_text=True)
    assert 'test-secret' not in client.get('/api/v1/llm/config').get_data(as_text=True)
    assert client.patch('/api/v1/llm/config',json={'base_url':'https://user:secret@example.test/v1'}).status_code==400
    def unauthorized(*args,**kwargs):raise urllib.error.HTTPError('unused',401,'SECRET PROVIDER BODY',{},None)
    monkeypatch.setattr('urllib.request.urlopen',unauthorized)
    r=client.post('/api/v1/llm/test');assert r.status_code==502 and 'SECRET' not in r.get_data(as_text=True)
    def timeout(*args,**kwargs):raise TimeoutError('SECRET REQUEST')
    monkeypatch.setattr('urllib.request.urlopen',timeout)
    r=client.post('/api/v1/llm/test');assert r.status_code==504 and 'SECRET' not in r.get_data(as_text=True)


def test_canonical_merge_aggregates_outputs_without_rewriting_original_question_ids(client,app):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    child=question(client);root=question(client)
    generated=generate(client,fake,child)
    saved=client.post(f"/api/v1/assistant/previews/{generated['preview_id']}/save",json={'save_kind':'output'}).get_json()['output']
    relation=client.get(f'/api/v1/questions/{root}/similar-candidates').get_json()['candidates'][0]
    reviewed=client.patch(f"/api/v1/question-relations/{relation['id']}",json={'question_id':root,'relation_type':'same_question','decision_status':'accepted','expected_review_token':relation['review_token']})
    assert reviewed.status_code==200
    preview=client.get(f'/api/v1/questions/{root}/merge-preview?source_question_id={child}').get_json()
    merged=client.post(f'/api/v1/questions/{root}/merge',json={'canonical_id':root,'source_question_id':child,'relation_id':relation['id'],'preview_token':preview['preview_token'],'topic_ids':[],'tag_ids':[]})
    assert merged.status_code==200
    rows=client.get(f'/api/v1/questions/{root}/assistant-outputs').get_json()
    assert len(rows)==1 and rows[0]['question_id']==child and rows[0]['id']==saved['id']
    history=client.get(f'/api/v1/questions/{root}/history').get_json()
    assert history['assistant_outputs']==rows


def test_explicit_exclusions_and_archived_materials_do_not_send_content(client,app):
    fake=FakeProvider();app.config['LLM_PROVIDER_FACTORY']=lambda _:fake
    qid=question(client);p=project(client)
    excluded=upload(client,content='EXCLUDED MATERIAL BODY',project_id=p['id'])
    archived=upload(client,content='ARCHIVED MATERIAL BODY',project_id=p['id'])
    client.post(f"/api/v1/materials/{archived['id']}/archive")
    result=generate(client,fake,qid,context={'project_ids':[p['id']],'exclude_material_ids':[excluded['id']]})
    sent=json.dumps(fake.messages[-1])
    assert 'EXCLUDED MATERIAL BODY' not in sent and 'ARCHIVED MATERIAL BODY' not in sent
    assert {s['material_id'] for s in result['sources']}=={p['profile_material_id']}
